from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# --- Visual Identity: Ultrathink ---
KOREAN_FONT = "Apple SD Gothic Neo"
ENGLISH_FONT = "Helvetica Neue"

COLOR_BG = RGBColor(15, 23, 42)      # Slate 900 (Darker)
COLOR_TEXT_MAIN = RGBColor(248, 250, 252) # Slate 50
COLOR_TEXT_SUB = RGBColor(148, 163, 184)  # Slate 400
COLOR_ACCENT_PURPLE = RGBColor(192, 132, 252) # Purple 400
COLOR_ACCENT_GREEN = RGBColor(134, 239, 172)  # Green 300
COLOR_ACCENT_GOLD = RGBColor(250, 204, 21)    # Yellow 400
COLOR_ACCENT_BLUE = RGBColor(96, 165, 250)    # Blue 400

IMAGE_DIR = "/Users/ted/vivid/scripts/ppt_screenshots"

def set_font(run, font_name, size, color, bold=False):
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run._r.get_or_add_rPr().set(qn('w:eastAsia'), font_name)

def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_image_safe(slide, img_name, left, top, width=None, height=None):
    img_path = os.path.join(IMAGE_DIR, img_name)
    if os.path.exists(img_path):
        try:
            if width and height:
                slide.shapes.add_picture(img_path, left, top, width=width, height=height)
            elif width:
                slide.shapes.add_picture(img_path, left, top, width=width)
            elif height:
                slide.shapes.add_picture(img_path, left, top, height=height)
            else:
                slide.shapes.add_picture(img_path, left, top)
        except Exception as e:
            print(f"Error: {img_name}: {e}")

def add_title(slide, text):
    # Minimalist Title
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(10), Inches(0.8))
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    set_font(p.runs[0], KOREAN_FONT, Pt(40), COLOR_TEXT_MAIN, bold=True)
    
    # Decorative Line
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.3), Inches(0.8), Inches(0.05))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_ACCENT_PURPLE
    shape.line.fill.background() # No line

# PPT Setup
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ===== SLIDE 1: Cover =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)

# Title Group
txBox = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11), Inches(2))
p = txBox.text_frame.paragraphs[0]
p.text = "AI 영상 제작 마스터 클래스"
set_font(p.runs[0], KOREAN_FONT, Pt(64), COLOR_ACCENT_PURPLE, bold=True)

# Subtitle
txBox = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "인간이 가장 잘하는 일에 집중하기"
set_font(p.runs[0], KOREAN_FONT, Pt(32), COLOR_TEXT_MAIN)

# Footer
txBox = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "Crebit Dimension × Arkain Studio"
set_font(p.runs[0], ENGLISH_FONT, Pt(20), COLOR_TEXT_SUB)

# ===== SLIDE 2: Vision (Why?) =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "왜 이 과정인가?")

# Left: Text Block
txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(4))
tf = txBox.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "바이럴 조회수나\n단기 수익화 교육이 아닙니다."
set_font(p.runs[0], KOREAN_FONT, Pt(28), COLOR_TEXT_SUB)
p.space_after = Pt(40)

p = tf.add_paragraph()
p.text = "AI가 아무리 발전해도,\n결정의 주체는 인간입니다."
set_font(p.runs[0], KOREAN_FONT, Pt(32), COLOR_TEXT_MAIN, bold=True)

# Right: Emphasis
txBox = slide.shapes.add_textbox(Inches(7.5), Inches(2.5), Inches(5), Inches(3))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = '"이 장면이 왜 아름다운가?"'
set_font(p.runs[0], KOREAN_FONT, Pt(36), COLOR_ACCENT_PURPLE, bold=True)
p.space_after = Pt(20)

p = tf.add_paragraph()
p.text = '"왜 가슴에 남는가?"'
set_font(p.runs[0], KOREAN_FONT, Pt(36), COLOR_ACCENT_GREEN, bold=True)

# Bottom Highlight
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3), Inches(6), Inches(7.33), Inches(0.8))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(30, 41, 59) # Slate 800
shape.line.color.rgb = COLOR_ACCENT_GOLD
txBox = slide.shapes.add_textbox(Inches(3), Inches(6.15), Inches(7.33), Inches(0.8))
p = txBox.text_frame.paragraphs[0]
p.text = "Human-in-the-Loop: 미학적 판단과 서사적 감각"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(22), COLOR_ACCENT_GOLD, bold=True)


# ===== SLIDE 3: Data-Driven RAG =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "거장의 미학을 데이터로")

items = [
    ("Masterpiece DNA", "왕가위, 타르코프스키\n거장 작품 RAG 해석", COLOR_ACCENT_PURPLE),
    ("Visual Analysis", "색감, 구도, 리듬, 앵글\n정량적 미학 데이터", COLOR_ACCENT_BLUE),
    ("Sound Logic", "음악과 영상의 조화\n논문 기반 사운드 설계", COLOR_ACCENT_GREEN),
    ("Consistent Flow", "긴 영상의 시각적 일관성\n동적 프롬프트 시스템", COLOR_ACCENT_GOLD)
]

# 2x2 Grid with modern cards
start_x, start_y = Inches(1.5), Inches(2.0)
gap_x, gap_y = Inches(0.5), Inches(0.5)
card_w, card_h = Inches(5.0), Inches(2.2)

for i, (title, desc, color) in enumerate(items):
    row, col = divmod(i, 2)
    left = start_x + col * (card_w + gap_x)
    top = start_y + row * (card_h + gap_y)
    
    # Card Background
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 41, 59)
    shape.line.color.rgb = color # Border accent
    shape.line.width = Pt(1.5)
    
    # Text
    txBox = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.3), card_w - Inches(0.6), card_h - Inches(0.6))
    p = txBox.text_frame.paragraphs[0]
    p.text = title
    set_font(p.runs[0], ENGLISH_FONT, Pt(20), color, bold=True)
    
    p = txBox.text_frame.add_paragraph()
    p.text = ""
    p = txBox.text_frame.add_paragraph()
    p.text = desc
    set_font(p.runs[0], KOREAN_FONT, Pt(18), COLOR_TEXT_MAIN)


# ===== SLIDE 4: Workflow Overview =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
# Center Title
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.33), Inches(0.8))
p = txBox.text_frame.paragraphs[0]
p.text = "4단계 워크플로우"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(36), COLOR_TEXT_MAIN, bold=True)

# Image Floating in Center
add_image_safe(slide, "flow-workflow.png", Inches(1.5), Inches(1.5), width=Inches(10.33))

# Bottom Tagline
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(12.33), Inches(0.8))
p = txBox.text_frame.paragraphs[0]
p.text = "1주차부터 '만들면서 배우는' 실전형 프로세스"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(22), COLOR_TEXT_SUB)


# ===== SLIDE 5: Stage 1 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "1. Identity & Planning")

# Apps Row
apps = [
    ("심연의 거울", "abyss-mirror.png", "창작 DNA 발견"),
    ("레퍼런스 해석기", "reference-decoder.png", "연출 분석"),
    ("시나리오 생성기", "story-architect.png", "스토리 구조화")
]

start_x = Inches(1)
card_w = Inches(3.5)
gap = Inches(0.6)

for i, (name, img, desc) in enumerate(apps):
    left = start_x + i * (card_w + gap)
    top = Inches(2.0)
    
    # Image
    add_image_safe(slide, img, left, top, width=card_w)
    
    # Text Group
    txBox = slide.shapes.add_textbox(left, top + Inches(2.2), card_w, Inches(1.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(20), COLOR_ACCENT_PURPLE, bold=True)
    
    p = txBox.text_frame.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(16), COLOR_TEXT_SUB)


# ===== SLIDE 6: Stage 2 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "2. Pre-production")

apps = [
    ("미학 디렉터", "aesthetic-director.png", "스타일 가이드"),
    ("스토리보드", "storyboard-sketch.png", "시각적 설계"),
    ("사운드 크래프터", "sound-crafter.png", "BGM/SFX"),
    ("프롬프트 연금술", "prompt-alchemy.png", "Lang 변환")
]

start_x = Inches(0.5)
card_w = Inches(2.8)
gap = Inches(0.4)

for i, (name, img, desc) in enumerate(apps):
    left = start_x + i * (card_w + gap)
    top = Inches(2.0)
    add_image_safe(slide, img, left, top, width=card_w)
    
    txBox = slide.shapes.add_textbox(left, top + Inches(1.8), card_w, Inches(1.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(18), COLOR_ACCENT_BLUE, bold=True)
    
    p = txBox.text_frame.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(14), COLOR_TEXT_SUB)


# ===== SLIDE 7: Stage 3 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "3. Production")

apps = [
    ("비주얼 리얼라이저", "visual-realizer.png", "Keyframe Generation"),
    ("비디오 메이커", "video-maker.png", "Video Generation")
]

start_x = Inches(1.5)
card_w = Inches(4.8)
gap = Inches(0.8)

for i, (name, img, desc) in enumerate(apps):
    left = start_x + i * (card_w + gap)
    top = Inches(2.0)
    add_image_safe(slide, img, left, top, width=card_w)
    
    txBox = slide.shapes.add_textbox(left, top + Inches(3.0), card_w, Inches(1.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(22), COLOR_ACCENT_GREEN, bold=True)
    
    p = txBox.text_frame.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(16), COLOR_TEXT_SUB)


# ===== SLIDE 8: Stage 4 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "4. Finishing")

add_image_safe(slide, "quality-director.png", Inches(1.5), Inches(1.8), width=Inches(10.33))

txBox = slide.shapes.add_textbox(Inches(1.5), Inches(6.0), Inches(10.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "퀄리티 디렉터: 전문가급 품질 검수 및 업스케일링"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(24), COLOR_ACCENT_GOLD, bold=True)


# ===== SLIDE 9: Dimension Merge =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "차원 확장: 아이디어의 성장")

add_image_safe(slide, "flow-train-view.png", Inches(0.5), Inches(1.5), width=Inches(12.33))

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.33), Inches(0.8))
p = txBox.text_frame.paragraphs[0]
p.text = "단순한 아이디어가 각 차원을 거치며 완성된 작품으로 진화합니다."
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(20), COLOR_TEXT_SUB)


# ===== SLIDE 10: Outcomes =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "What You Make")

items = [
    ("3분 애니메이션 MV", COLOR_ACCENT_PURPLE),
    ("3분 수준급 숏필름", COLOR_ACCENT_GREEN),
    ("AI OTT 레벨 시네마", COLOR_ACCENT_GOLD)
]

start_y = Inches(2.2)
gap = Inches(1.4)

for i, (text, color) in enumerate(items):
    top = start_y + i * gap
    # Minimalist underline style
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), top + Inches(0.8), Inches(9.33), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    
    txBox = slide.shapes.add_textbox(Inches(2), top, Inches(9.33), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    set_font(p.runs[0], KOREAN_FONT, Pt(40), COLOR_TEXT_MAIN, bold=True)


# ===== SLIDE 11: Market Context =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "Market Direction")

# Stats Big Number
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.33), Inches(2))
p = txBox.text_frame.paragraphs[0]
p.text = "$11,000,000,000"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], ENGLISH_FONT, Pt(80), COLOR_ACCENT_PURPLE, bold=True)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "2025 글로벌 숏폼 드라마 시장 규모 (Sensor Tower)"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(20), COLOR_TEXT_SUB)

# Context Text
txBox = slide.shapes.add_textbox(Inches(2), Inches(5.0), Inches(9.33), Inches(2))
p = txBox.text_frame.paragraphs[0]
p.text = "바이럴 시장은 포화되었습니다.\n이제는 AI OTT 오리지널 콘텐츠의 시대입니다."
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(24), COLOR_TEXT_MAIN)


# ===== SLIDE 12: ATC Producer =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "ATC Producer")

# Concept Definition
txBox = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11), Inches(1.5))
p = txBox.text_frame.paragraphs[0]
p.text = "AI-to-Content Producer"
set_font(p.runs[0], ENGLISH_FONT, Pt(48), COLOR_ACCENT_GOLD, bold=True)

# Bullet points
items = [
    "AI 도구로 프로덕션급 영상을 기획·제작",
    "3분~장편을 OTT 플랫폼에 공급",
    "기술(Tech)과 미학(Art)을 잇는 전문가"
]

start_y = Inches(3.5)
for i, item in enumerate(items):
    top = start_y + i * Inches(0.8)
    # Bullet dot
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), top + Inches(0.15), Inches(0.15), Inches(0.15))
    dot.fill.solid()
    dot.fill.fore_color.rgb = COLOR_ACCENT_GOLD
    dot.line.fill.background()
    
    txBox = slide.shapes.add_textbox(Inches(1.4), top, Inches(10), Inches(0.6))
    p = txBox.text_frame.paragraphs[0]
    p.text = item
    set_font(p.runs[0], KOREAN_FONT, Pt(24), COLOR_TEXT_MAIN)


# ===== SLIDE 13: Internship =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "Career Path")

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "파트너사 3개월 인턴십 우선 연계"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(28), COLOR_TEXT_MAIN, bold=True)

# Logos/Names Grid
companies = ["Arkain Studio", "M83 Studio", "Spoonlabs", "+ 4 Partners"]
colors = [COLOR_ACCENT_PURPLE, COLOR_ACCENT_BLUE, COLOR_ACCENT_GREEN, COLOR_TEXT_SUB]

start_x = Inches(1.5)
gap = Inches(0.5)
w = Inches(2.2) # (13.3 - 3 - 1.5)/4 approx 
w = Inches(2.4)

for i, (company, color) in enumerate(zip(companies, colors)):
    left = start_x + i * (w + gap)
    top = Inches(3.5)
    
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 41, 59)
    shape.line.color.rgb = color
    shape.line.width = Pt(1)
    
    txBox = slide.shapes.add_textbox(left, top + Inches(0.5), w, Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = company
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], ENGLISH_FONT, Pt(18), color, bold=True)

# Bottom
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.0), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "단순 수료가 아닌, 실제 프로젝트 투입"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(20), COLOR_ACCENT_GOLD)


# ===== SLIDE 14: Closing =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)

# Quote Design
txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(3))
p = txBox.text_frame.paragraphs[0]
p.text = '"AI가 90%를 해주는 시대,'
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(36), COLOR_TEXT_SUB)

p = txBox.text_frame.add_paragraph()
p.text = '당신은 가장 중요한 10%에 집중하세요."'
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(44), COLOR_ACCENT_PURPLE, bold=True)

# Save
output_path = "/Users/ted/Desktop/AI_영상마스터_클래스_2025_Ultrathink.pptx"
prs.save(output_path)
print(f"완료: {output_path}")
