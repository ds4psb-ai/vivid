from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

KOREAN_FONT = "Apple SD Gothic Neo"
IMAGE_DIR = "/Users/ted/vivid/scripts/ppt_screenshots"

# 슬라이드 헬퍼 함수
def set_font(run, font_name, size, color, bold=False):
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run._r.get_or_add_rPr().set(qn('w:eastAsia'), font_name)

def set_background(slide, color=RGBColor(10, 10, 10)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_image_safe(slide, img_name, left, top, width=None, height=None):
    img_path = os.path.join(IMAGE_DIR, img_name)
    if os.path.exists(img_path):
        try:
            if width and height:
                slide.shapes.add_picture(img_path, left, top, width=width, height=height)
            elif width:
                slide.shapes.add_picture(img_path, left, top, width=width)
            elif height:
                slide.shapes.add_picture(img_path, left, top, height=height)
            else:
                slide.shapes.add_picture(img_path, left, top)
        except Exception as e:
            print(f"Error: {img_name}: {e}")

def add_title(slide, text, top=Inches(0.3), size=Pt(36), color=RGBColor(255, 255, 255)):
    txBox = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.33), Inches(0.8))
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, size, color, bold=True)

# PPT 생성
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ===== SLIDE 1: 표지 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(12.33), Inches(1.5))
p = txBox.text_frame.paragraphs[0]
p.text = "AI 영상 제작 마스터 클래스"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(54), RGBColor(168, 85, 247), bold=True)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "인간이 가장 잘하는 일에 집중하기"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(28), RGBColor(226, 232, 240))

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "Crebit Dimension × Arkain Studio"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(100, 100, 100))

# ===== SLIDE 2: 왜 이 과정인가? (목적) =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "왜 이 과정인가?")

# 메인 메시지
txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "이 과정은 바이럴 조회수나 단기 수익화를 위한 교육이 아닙니다."
set_font(p.runs[0], KOREAN_FONT, Pt(24), RGBColor(200, 200, 200))

p = tf.add_paragraph()
p.text = ""
p = tf.add_paragraph()
p.text = "AI가 아무리 발전해도"
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(150, 150, 150))

p = tf.add_paragraph()
p.text = '"이 장면이 왜 아름다운가"'
set_font(p.runs[0], KOREAN_FONT, Pt(28), RGBColor(168, 85, 247), bold=True)

p = tf.add_paragraph()
p.text = '"이 전개가 왜 가슴에 남는가"'
set_font(p.runs[0], KOREAN_FONT, Pt(28), RGBColor(134, 239, 172), bold=True)

p = tf.add_paragraph()
p.text = ""
p = tf.add_paragraph()
p.text = "를 결정하는 건 여전히 인간의 몫입니다."
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(150, 150, 150))

# 하단 강조
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.3), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "Human-in-the-Loop: 미학적 판단과 서사적 감각에 집중"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(22), RGBColor(255, 215, 0), bold=True)

# ===== SLIDE 3: RAG 데이터 기반 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "거장의 미학을 데이터로")

# 4개 아이콘 그리드
items = [
    ("거장 작품 해석", "왕가위, 타르코프스키,\n아케인 시리즈", RGBColor(168, 85, 247)),
    ("마스터피스 분석", "색감, 구도, 리듬\nRAG 데이터셋", RGBColor(59, 130, 246)),
    ("논문 기반 사운드", "음악과 영상의 조화\n사운드 설계 가이드", RGBColor(34, 197, 94)),
    ("동적 일관성", "긴 영상에서도\n시각적 일관성 유지", RGBColor(234, 179, 8))
]

for i, (title, desc, color) in enumerate(items):
    left = Inches(0.5 + (i * 3.2))
    
    # 박스
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2), Inches(3), Inches(4))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 30, 40)
    shape.line.color.rgb = color
    shape.line.width = Pt(2)
    
    # 텍스트
    txBox = slide.shapes.add_textbox(left, Inches(2.5), Inches(3), Inches(3.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(22), color, bold=True)
    
    p = tf.add_paragraph()
    p.text = ""
    p = tf.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(16), RGBColor(180, 180, 180))

# ===== SLIDE 4: 4단계 워크플로우 개요 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "4단계 워크플로우")

add_image_safe(slide, "flow-workflow.png", Inches(1.5), Inches(1.5), width=Inches(10.33))

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.3), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "1주차부터 '만들면서 배우는' 방식으로 진입"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(22), RGBColor(200, 200, 200))

# ===== SLIDE 5: 1단계 기획 (3 Apps) =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "1단계: 정체성 & 기획")

apps = [
    ("심연의 거울", "abyss-mirror.png", "창작 DNA 발견"),
    ("레퍼런스 해석기", "reference-decoder.png", "연출 기법 분석"),
    ("시나리오 생성기", "story-architect.png", "구조화된 스토리")
]
for i, (name, img, desc) in enumerate(apps):
    left = Inches(0.5 + (i * 4.1))
    add_image_safe(slide, img, left, Inches(1.5), width=Inches(3.8))
    txBox = slide.shapes.add_textbox(left, Inches(5.2), Inches(3.8), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(168, 85, 247), bold=True)
    p = tf.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(16), RGBColor(180, 180, 180))

# ===== SLIDE 6: 2단계 사전제작 (4 Apps) =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "2단계: 사전 제작")

apps = [
    ("미학 디렉터", "aesthetic-director.png", "스타일 가이드"),
    ("스토리보드", "storyboard-sketch.png", "시각적 컷 설계"),
    ("사운드 크래프터", "sound-crafter.png", "BGM/효과음"),
    ("프롬프트 연금술", "prompt-alchemy.png", "AI 언어 변환")
]
for i, (name, img, desc) in enumerate(apps):
    left = Inches(0.3 + (i * 3.2))
    add_image_safe(slide, img, left, Inches(1.5), width=Inches(3))
    txBox = slide.shapes.add_textbox(left, Inches(5.2), Inches(3), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(18), RGBColor(59, 130, 246), bold=True)
    p = tf.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(14), RGBColor(180, 180, 180))

# ===== SLIDE 7: 3단계 제작 (2 Apps) =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "3단계: 제작")

apps = [
    ("비주얼 리얼라이저", "visual-realizer.png", "고품질 키프레임 생성"),
    ("비디오 메이커", "video-maker.png", "영상 생성 & 모션 제어")
]
for i, (name, img, desc) in enumerate(apps):
    left = Inches(1 + (i * 6))
    add_image_safe(slide, img, left, Inches(1.5), width=Inches(5.5))
    txBox = slide.shapes.add_textbox(left, Inches(5.5), Inches(5.5), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(22), RGBColor(34, 197, 94), bold=True)
    p = tf.add_paragraph()
    p.text = desc
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(16), RGBColor(180, 180, 180))

# ===== SLIDE 8: 4단계 완성 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "4단계: 완성")

add_image_safe(slide, "quality-director.png", Inches(2), Inches(1.5), width=Inches(9.33))

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(12.33), Inches(1.5))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "퀄리티 디렉터: 전문가 품질 검수"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(24), RGBColor(234, 179, 8), bold=True)

p = tf.add_paragraph()
p.text = "혼자 만들어도 프로덕션 퀄리티 보장"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(18), RGBColor(200, 200, 200))

# ===== SLIDE 9: 차원 확장 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "차원 확장: 아이디어의 성장")

add_image_safe(slide, "flow-train-view.png", Inches(1), Inches(1.5), width=Inches(11.33))

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.33), Inches(0.8))
p = txBox.text_frame.paragraphs[0]
p.text = "단순한 아이디어가 각 차원을 거치며 구체적인 결과물로 성장"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(200, 200, 200))

# ===== SLIDE 10: 결과물 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "코스 종료 시 결과물")

items = [
    ("3분 애니메이션 MV", RGBColor(168, 85, 247)),
    ("3분 수준급 숏필름", RGBColor(134, 239, 172)),
    ("AI OTT 레벨 작품", RGBColor(234, 179, 8))
]

for i, (text, color) in enumerate(items):
    top = Inches(2 + (i * 1.5))
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3), top, Inches(7.33), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 30, 40)
    shape.line.color.rgb = color
    shape.line.width = Pt(2)
    
    txBox = slide.shapes.add_textbox(Inches(3), top, Inches(7.33), Inches(1.2))
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(28), color, bold=True)

# ===== SLIDE 11: 방향 - AI OTT 시장 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "방향: AI OTT 시장")

txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(5.5))
tf = txBox.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "바이럴 콘텐츠 시장은 2024년 반짝 성장 후 포화 상태"
set_font(p.runs[0], KOREAN_FONT, Pt(22), RGBColor(150, 150, 150))

p = tf.add_paragraph()
p.text = ""
p = tf.add_paragraph()
p.text = "2025년 글로벌 숏폼 드라마 시장"
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(200, 200, 200))

p = tf.add_paragraph()
p.text = "110억 달러"
set_font(p.runs[0], KOREAN_FONT, Pt(48), RGBColor(168, 85, 247), bold=True)

p = tf.add_paragraph()
p.text = "(Sensor Tower 전망)"
set_font(p.runs[0], KOREAN_FONT, Pt(16), RGBColor(100, 100, 100))

p = tf.add_paragraph()
p.text = ""
p = tf.add_paragraph()
p.text = "아케인 스튜디오는 이미 AI 숏드라마 제작 계약 체결"
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(134, 239, 172))

# ===== SLIDE 12: ATC 프로듀서 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "ATC 프로듀서란?")

items = [
    "AI 도구를 활용해 프로덕션급 영상을 기획·제작하는 전문가",
    "3분~장편 작품을 AI OTT 플랫폼에 공급하는 크리에이터",
    "기술과 미학을 모두 이해하는 Human-in-the-Loop 핵심 인력"
]

txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(4))
tf = txBox.text_frame
tf.word_wrap = True

for i, item in enumerate(items):
    if i == 0:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.text = f"• {item}"
    set_font(p.runs[0], KOREAN_FONT, Pt(22), RGBColor(200, 200, 200))
    p.space_after = Pt(30)

# 하단 강조
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "ATC = AI-to-Content"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(24), RGBColor(234, 179, 8), bold=True)

# ===== SLIDE 13: 인턴십 연계 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "인턴십 및 취업 연계")

txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "우수 수료자는 파트너사 3개월 인턴십 프로그램 우선 연결"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(22), RGBColor(200, 200, 200))

# 파트너사 리스트
companies = ["아케인 스튜디오", "M83 스튜디오", "스푼랩스", "+ 4개 파트너사"]
for i, company in enumerate(companies):
    left = Inches(0.5 + (i * 3.2))
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(4), Inches(3), Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(30, 30, 40)
    shape.line.color.rgb = RGBColor(168, 85, 247)
    shape.line.width = Pt(2)
    
    txBox = slide.shapes.add_textbox(left, Inches(4.3), Inches(3), Inches(1))
    p = txBox.text_frame.paragraphs[0]
    p.text = company
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], KOREAN_FONT, Pt(18), RGBColor(200, 200, 200), bold=True)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.3), Inches(12.33), Inches(1))
p = txBox.text_frame.paragraphs[0]
p.text = "단순 수료가 아닌, 실제 프로젝트 투입"
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(20), RGBColor(134, 239, 172))

# ===== SLIDE 14: 마무리 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(3), Inches(12.33), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = '"AI가 90%를 해주는 시대,'
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(32), RGBColor(200, 200, 200))

p = tf.add_paragraph()
p.text = '당신은 가장 중요한 10%에 집중하세요."'
p.alignment = PP_ALIGN.CENTER
set_font(p.runs[0], KOREAN_FONT, Pt(32), RGBColor(168, 85, 247), bold=True)

output_path = "/Users/ted/Desktop/AI_영상마스터_클래스_2025.pptx"
prs.save(output_path)
print(f"완료: {output_path}")
print(f"총 슬라이드: {len(prs.slides)}장")
