from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 제목
title = doc.add_heading('넥스트러너스 협업 제안서', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('미팅 주제: 코딩+콘텐츠 = 전인적 사업가 양성')
doc.add_paragraph('대상: 이한별 대표 (넥스트러너스, 오즈코딩스쿨, 라이프해킹스쿨)')
doc.add_paragraph()

# 1. 협업 배경
doc.add_heading('1. 협업 배경', level=1)
doc.add_paragraph('넥스트러너스와 크레빗은 같은 방향을 바라보고 있다.')

table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
table.rows[0].cells[0].text = '넥스트러너스 (오즈코딩스쿨)'
table.rows[0].cells[1].text = '크레빗 (AI 영상 마스터)'
table.rows[1].cells[0].text = '바이브코딩'
table.rows[1].cells[1].text = '컨셉 디렉팅'
table.rows[2].cells[0].text = '개발자가 AI 코드 검증'
table.rows[2].cells[1].text = '디렉터가 AI 영상 검증'
table.rows[3].cells[0].text = '결과물: 1인 SaaS/서비스'
table.rows[3].cells[1].text = '결과물: 브랜드 스토리, 마케팅 영상'
table.rows[4].cells[0].text = '제품을 만든다'
table.rows[4].cells[1].text = '제품을 판다'

doc.add_paragraph()
doc.add_paragraph('공통 목표: 혼자서 서비스도 만들고 마케팅도 할 수 있는 자립형 사업가 양성')
doc.add_paragraph()

# 2. 크레빗 영상 제작 파이프라인
doc.add_heading('2. 크레빗 영상 제작 파이프라인 (5단계, 8개 앱)', level=1)
doc.add_paragraph('도구 사용법이 아니라 프로덕션 워크플로우를 가르친다.')

workflow = doc.add_paragraph()
workflow.add_run('1단계 정체성: ').bold = True
workflow.add_run('심연의 거울 → 레퍼런스 해석기\n')
workflow.add_run('2단계 기획: ').bold = True
workflow.add_run('시나리오 생성기 → 사운드 디자이너\n')
workflow.add_run('3단계 설계: ').bold = True
workflow.add_run('스토리보드 스케치 → 첫 프레임 생성기\n')
workflow.add_run('4단계 제작: ').bold = True
workflow.add_run('비디오 메이커\n')
workflow.add_run('5단계 품질: ').bold = True
workflow.add_run('퀄리티 디렉터')

doc.add_paragraph()

apps = [
    '1. 심연의 거울: 창업자의 취향/DNA를 분석 (RAG 기반)',
    '2. 레퍼런스 해석기: "이런 느낌으로 만들어줘" 분석 (멀티모달)',
    '3. 시나리오 생성기: 아이디어를 판매 가능한 스토리로 변환',
    '4. 사운드 디자이너: 브랜드 무드에 맞는 BGM/효과음 생성',
    '5. 스토리보드 스케치: 시각 기획 및 프롬프트 튜닝',
    '6. 첫 프레임 생성기: 핵심 비주얼 이미지 생성',
    '7. 비디오 메이커: 이미지를 영상으로 확장 (Veo, Kling, Runway)',
    '8. 퀄리티 디렉터: 자동 품질 검수 및 업스케일링'
]
for app in apps:
    doc.add_paragraph(app)

doc.add_paragraph()

# 3. 협업 모델
doc.add_heading('3. 협업 모델: 퍼널 전략', level=1)
doc.add_paragraph('넥스트러너스가 온라인 기초 과정(국비)을 운영하고, 상위 10%를 크레빗 오프라인 마스터 과정으로 연결한다.')

model_table = doc.add_table(rows=3, cols=2)
model_table.style = 'Table Grid'
model_table.rows[0].cells[0].text = '온라인 기초 (대량)'
model_table.rows[0].cells[1].text = '운영: 넥스트러너스 (국비)\n커리큘럼: 크레빗 제공\n수익: 국비 전액 넥스트러너스'
model_table.rows[1].cells[0].text = '선별'
model_table.rows[1].cells[1].text = '쇼케이스 평가 → 상위 10% 선발'
model_table.rows[2].cells[0].text = '오프라인 마스터 (정예)'
model_table.rows[2].cells[1].text = '운영: 크레빗/아케인\n커리큘럼: 퀄리티 디렉팅, 상업 영상\n수익: 7:3 쉐어'

doc.add_paragraph()

# 4. 이한별 대표에게 할 말
doc.add_heading('4. 이한별 대표에게 할 말', level=1)

pitch1 = doc.add_paragraph()
pitch1.add_run('첫째, 1인 사업가 커리큘럼 완성\n').bold = True
pitch1.add_run('"대표님은 제품을 만드는 개발자를 양성하고 계십니다. 그런데 그 제품을 팔 수 있나요? 우리 과정은 제품 엔진에 마케팅 엔진을 더합니다."')

pitch2 = doc.add_paragraph()
pitch2.add_run('둘째, 수강생 품질 관리\n').bold = True
pitch2.add_run('"마스터 클래스라는 명확한 목표를 제시하면 수강생 동기부여가 됩니다."')

pitch3 = doc.add_paragraph()
pitch3.add_run('셋째, 기술 시너지\n').bold = True
pitch3.add_run('"백엔드 RAG와 저희 RAG(심연의 거울)는 같은 기술입니다. 도메인만 다릅니다."')

pitch4 = doc.add_paragraph()
pitch4.add_run('넷째, 졸업작품 영상화\n').bold = True
pitch4.add_run('"오즈코딩스쿨 졸업생의 서비스 런칭 영상을 크레빗으로 제작하는 모듈을 함께 설계할 수 있습니다."')

doc.add_paragraph()

# 5. 미팅 준비물
doc.add_heading('5. 미팅 준비물', level=1)
doc.add_paragraph('1. 크레빗 플랫폼 라이브 데모 (심연의 거울, 30초 광고 생성)')
doc.add_paragraph('2. 수강생 결과물 3개 이상 (뮤직비디오, 숏드라마, 광고)')
doc.add_paragraph('3. 국비지원 과정 제안서 초안')

doc.add_paragraph()

# 6. 질문할 것
doc.add_heading('6. 질문할 것', level=1)
doc.add_paragraph('1. 2026년 K-디지털 트레이닝 신규 과정 계획이 있으신가요?')
doc.add_paragraph('2. 원더스랩 협력에서 AI 콘텐츠 영역은 어떻게 되나요?')
doc.add_paragraph('3. 오즈코딩스쿨 수강생 중 콘텐츠 제작 니즈가 있는 비중은?')
doc.add_paragraph('4. 라이프해킹스쿨 창업 과정에서 영상 마케팅은 어떻게 다루시나요?')

doc.add_paragraph()

# 핵심 한 줄
doc.add_heading('핵심 한 줄', level=1)
summary = doc.add_paragraph()
summary.add_run('바이브코딩이 개발을 민주화했듯이, 크레빗은 영상 제작을 민주화한다.\n')
summary.add_run('넥스트러너스의 창업가 양성 철학과 연결하면, 100만 1인 창업가가 각자의 브랜드 영상을 갖게 된다.')

# 저장
file_path = '/Users/ted/Desktop/넥스트러너스_협업제안서.docx'
doc.save(file_path)
print(f'완료: {file_path}')
